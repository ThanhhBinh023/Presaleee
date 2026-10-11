"""
Comprehensive Generator for 10 Enterprise Leads and 10 Full Proposals
Each Lead: Authentic Client RFP format (Sheet 'Request')
Each Proposal: 10 Standard Sheets:
  1. Overview
  2. Project Roles
  3. Project Management Approach
  4. RAID
  5. Acceptance Docs
  6. FR
  7. Tech-stack
  8. Non-FR
  9. RateCard
  10. Macro Plan & Quotation
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Typography & Palette
FONT_FAMILY = "Segoe UI"
COLOR_HEADER_BG = "0F172A"       # Slate 900
COLOR_HEADER_TEXT = "FFFFFF"     # White
COLOR_SECTION_BG = "1E3A8A"      # Blue 900
COLOR_SECTION_TEXT = "FFFFFF"    # White
COLOR_SUBSECTION_BG = "DBEAFE"   # Light Blue 100
COLOR_SUBSECTION_TEXT = "1E3A8A" # Dark Blue
COLOR_ZEBRA_BG = "F8FAFC"        # Slate 50
COLOR_BORDER = "CBD5E1"          # Slate 300

font_title = Font(name=FONT_FAMILY, size=15, bold=True, color="0F172A")
font_subtitle = Font(name=FONT_FAMILY, size=10, italic=True, color="475569")
font_section = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_SECTION_TEXT)
font_tbl_header = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_HEADER_TEXT)
font_cell = Font(name=FONT_FAMILY, size=9.5, color="1E293B")
font_cell_bold = Font(name=FONT_FAMILY, size=9.5, bold=True, color="1E293B")
font_meta_key = Font(name=FONT_FAMILY, size=9.5, bold=True, color="334155")
font_meta_val = Font(name=FONT_FAMILY, size=9.5, color="0F172A")

fill_header = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
fill_section = PatternFill(start_color=COLOR_SECTION_BG, end_color=COLOR_SECTION_BG, fill_type="solid")
fill_sub = PatternFill(start_color=COLOR_SUBSECTION_BG, end_color=COLOR_SUBSECTION_BG, fill_type="solid")
fill_zebra = PatternFill(start_color=COLOR_ZEBRA_BG, end_color=COLOR_ZEBRA_BG, fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)

header_border = Border(
    left=Side(style='thin', color="334155"),
    right=Side(style='thin', color="334155"),
    top=Side(style='medium', color="0F172A"),
    bottom=Side(style='medium', color="0F172A")
)

def auto_fit(ws, min_widths=None):
    ws.views.sheetView[0].showGridLines = True
    min_widths = min_widths or {}
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = cell.value
            if val is not None:
                lines = str(val).split("\n")
                line_max = max(len(l) for l in lines)
                if line_max > max_len:
                    max_len = line_max
        fitted = max(max_len + 3, min_widths.get(col_letter, 10))
        ws.column_dimensions[col_letter].width = min(fitted, 75)


def build_lead_workbook(file_path: str, rfp_meta: dict, rfp_sections: list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Request"

    ws["A1"] = rfp_meta["title"]
    ws["A1"].font = font_title
    ws.row_dimensions[1].height = 28

    ws["A2"] = f"From: {rfp_meta['client']}     |     Date sent: {rfp_meta['date']}     |     Contact: {rfp_meta['contact']}"
    ws["A2"].font = font_subtitle
    ws.row_dimensions[2].height = 20

    curr_r = 4
    for sec in rfp_sections:
        # Section header
        ws.merge_cells(start_row=curr_r, start_column=1, end_row=curr_r, end_column=2)
        c = ws.cell(row=curr_r, column=1, value=sec["section"])
        c.font = font_section
        c.fill = fill_section
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.cell(row=curr_r, column=2).border = thin_border
        ws.row_dimensions[curr_r].height = 24
        curr_r += 1

        for label, val in sec["items"]:
            cA = ws.cell(row=curr_r, column=1, value=label)
            cA.font = font_meta_key
            cA.fill = fill_zebra
            cA.border = thin_border
            cA.alignment = Alignment(horizontal="left", vertical="top")

            cB = ws.cell(row=curr_r, column=2, value=val)
            cB.font = font_meta_val
            cB.fill = fill_white
            cB.border = thin_border
            cB.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

            # Dynamic row height
            line_count = val.count("\n") + 1
            approx_height = max(24, line_count * 18, (len(val) // 90 + 1) * 16)
            ws.row_dimensions[curr_r].height = min(approx_height, 220)
            curr_r += 1

    ws.views.sheetView[0].showGridLines = True
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    wb.save(file_path)


def build_proposal_workbook(file_path: str, p_data: dict):
    wb = openpyxl.Workbook()

    # 1. Overview
    ws_ov = wb.active
    ws_ov.title = "Overview"
    ws_ov.merge_cells("A1:F1")
    ws_ov["A1"] = p_data["title"]
    ws_ov["A1"].font = font_title
    ws_ov.row_dimensions[1].height = 32

    ws_ov.merge_cells("A2:F2")
    ws_ov["A2"] = f"Client: {p_data['client']}  |  Proposal ID: {p_data['proposal_id']}  |  Date: {p_data['date']}  |  Validity: 30 Days"
    ws_ov["A2"].font = font_subtitle
    ws_ov.row_dimensions[2].height = 20

    r = 4
    for sec_name, rows in p_data["overview_sections"]:
        ws_ov.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = ws_ov.cell(row=r, column=1, value=sec_name)
        c.font = font_section
        c.fill = fill_section
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_ov.row_dimensions[r].height = 24
        r += 1

        for row_tuple in rows:
            if row_tuple[2] == "" and row_tuple[3] == "":
                ws_ov.cell(row=r, column=1, value=row_tuple[0]).font = font_meta_key
                ws_ov.cell(row=r, column=1).fill = fill_zebra
                ws_ov.cell(row=r, column=1).border = thin_border
                ws_ov.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="top")

                ws_ov.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
                val_c = ws_ov.cell(row=r, column=2, value=row_tuple[1])
                val_c.font = font_meta_val
                val_c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                for ci in range(2, 7):
                    ws_ov.cell(row=r, column=ci).border = thin_border
                line_count = row_tuple[1].count("\n") + 1
                ws_ov.row_dimensions[r].height = max(24, line_count * 18, (len(row_tuple[1]) // 75 + 1) * 16)
            else:
                ws_ov.cell(row=r, column=1, value=row_tuple[0]).font = font_meta_key
                ws_ov.cell(row=r, column=1).fill = fill_zebra
                ws_ov.cell(row=r, column=1).border = thin_border
                ws_ov.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")

                ws_ov.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
                ws_ov.cell(row=r, column=2, value=row_tuple[1]).font = font_meta_val
                ws_ov.cell(row=r, column=2).alignment = Alignment(horizontal="left", vertical="center")
                for ci in range(2, 4):
                    ws_ov.cell(row=r, column=ci).border = thin_border

                ws_ov.cell(row=r, column=4, value=row_tuple[2]).font = font_meta_key
                ws_ov.cell(row=r, column=4).fill = fill_zebra
                ws_ov.cell(row=r, column=4).border = thin_border
                ws_ov.cell(row=r, column=4).alignment = Alignment(horizontal="left", vertical="center")

                ws_ov.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
                ws_ov.cell(row=r, column=5, value=row_tuple[3]).font = font_meta_val
                ws_ov.cell(row=r, column=5).alignment = Alignment(horizontal="left", vertical="center")
                for ci in range(5, 7):
                    ws_ov.cell(row=r, column=ci).border = thin_border
                ws_ov.row_dimensions[r].height = 24
            r += 1
        r += 1
    auto_fit(ws_ov, {"A": 26, "B": 28, "C": 20, "D": 26, "E": 28, "F": 15})

    # 2. Project Roles
    ws_roles = wb.create_sheet(title="Project Roles")
    ws_roles.merge_cells("A1:G1")
    ws_roles["A1"] = "PROJECT ORGANIZATION & TEAM ROLES"
    ws_roles["A1"].font = font_title
    ws_roles.row_dimensions[1].height = 28
    ws_roles.merge_cells("A2:G2")
    ws_roles["A2"] = "Governance structure, resource allocation and RACI matrix."
    ws_roles["A2"].font = font_subtitle
    ws_roles.row_dimensions[2].height = 18

    headers_roles = ["Role Title", "Organization", "Allocation", "Primary Responsibilities", "Key Deliverables Owned", "Seniority"]
    ws_roles.row_dimensions[4].height = 24
    for i, h in enumerate(headers_roles, 1):
        cell = ws_roles.cell(row=4, column=i, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_idx, r_data in enumerate(p_data["roles_data"], 5):
        ws_roles.row_dimensions[r_idx].height = 28
        fill_curr = fill_zebra if r_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(r_data, 1):
            cell = ws_roles.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.fill = fill_curr
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [2, 3, 6] else "left", vertical="center", wrap_text=True)
    auto_fit(ws_roles, {"A": 26, "B": 15, "C": 18, "D": 45, "E": 35, "F": 14})

    # 3. Project Management Approach
    ws_pm = wb.create_sheet(title="Project Management Approach")
    ws_pm.merge_cells("A1:E1")
    ws_pm["A1"] = "PROJECT MANAGEMENT & DELIVERY METHODOLOGY"
    ws_pm["A1"].font = font_title
    ws_pm.row_dimensions[1].height = 28
    ws_pm.merge_cells("A2:E2")
    ws_pm["A2"] = "Agile Scrum framework, ceremonies, quality gates, and communication cadence."
    ws_pm["A2"].font = font_subtitle
    ws_pm.row_dimensions[2].height = 18
    r = 4
    for sec_t, items in p_data["pm_sections"]:
        ws_pm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c = ws_pm.cell(row=r, column=1, value=sec_t)
        c.font = font_section
        c.fill = fill_section
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_pm.row_dimensions[r].height = 24
        r += 1
        for label, val in items:
            ws_pm.cell(row=r, column=1, value=label).font = font_meta_key
            ws_pm.cell(row=r, column=1).fill = fill_zebra
            ws_pm.cell(row=r, column=1).border = thin_border
            ws_pm.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")
            ws_pm.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            val_cell = ws_pm.cell(row=r, column=2, value=val)
            val_cell.font = font_meta_val
            val_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            for ci in range(2, 6):
                ws_pm.cell(row=r, column=ci).border = thin_border
            ws_pm.row_dimensions[r].height = 24 if len(val) < 80 else 42
            r += 1
        r += 1
    auto_fit(ws_pm, {"A": 30, "B": 25, "C": 25, "D": 25, "E": 20})

    # 4. RAID
    ws_raid = wb.create_sheet(title="RAID")
    ws_raid.merge_cells("A1:G1")
    ws_raid["A1"] = "RAID LOG: RISKS, ASSUMPTIONS, ISSUES & DEPENDENCIES"
    ws_raid["A1"].font = font_title
    ws_raid.row_dimensions[1].height = 28
    ws_raid.merge_cells("A2:G2")
    ws_raid["A2"] = "Active risk management matrix monitored continuously."
    ws_raid["A2"].font = font_subtitle
    ws_raid.row_dimensions[2].height = 18
    headers_raid = ["ID", "Category", "Item Description", "Impact", "Probability", "Mitigation Strategy & Action Plan", "Owner"]
    ws_raid.row_dimensions[4].height = 24
    for i, h in enumerate(headers_raid, 1):
        cell = ws_raid.cell(row=4, column=i, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_idx, r_data in enumerate(p_data["raid_data"], 5):
        ws_raid.row_dimensions[r_idx].height = 32
        fill_curr = fill_zebra if r_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(r_data, 1):
            cell = ws_raid.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.fill = fill_curr
            cell.border = thin_border
            if c_idx in [1, 2, 4, 5, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    auto_fit(ws_raid, {"A": 12, "B": 15, "C": 45, "D": 12, "E": 12, "F": 45, "G": 20})

    # 5. Acceptance Docs
    ws_acc = wb.create_sheet(title="Acceptance Docs")
    ws_acc.merge_cells("A1:F1")
    ws_acc["A1"] = "QUALITY ASSURANCE, ACCEPTANCE CRITERIA & DEFECT SLA"
    ws_acc["A1"].font = font_title
    ws_acc.row_dimensions[1].height = 28
    ws_acc.merge_cells("A2:F2")
    ws_acc["A2"] = "Definition of Done (DoD) and production defect severity SLAs."
    ws_acc["A2"].font = font_subtitle
    ws_acc.row_dimensions[2].height = 18
    ws_acc.merge_cells("A4:F4")
    ws_acc["A4"] = "1. DEFINITION OF DONE (DoD) FOR SPRINT USER STORIES"
    ws_acc["A4"].font = font_section
    ws_acc["A4"].fill = fill_section
    ws_acc["A4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_acc.row_dimensions[4].height = 24
    r = 5
    for c_name, c_desc in p_data["dod_items"]:
        ws_acc.cell(row=r, column=1, value=c_name).font = font_meta_key
        ws_acc.cell(row=r, column=1).fill = fill_zebra
        ws_acc.cell(row=r, column=1).border = thin_border
        ws_acc.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws_acc.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws_acc.cell(row=r, column=2, value=c_desc).font = font_meta_val
        ws_acc.cell(row=r, column=2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for ci in range(2, 7):
            ws_acc.cell(row=r, column=ci).border = thin_border
        ws_acc.row_dimensions[r].height = 22
        r += 1
    r += 1
    ws_acc.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws_acc.cell(row=r, column=1, value="2. DEFECT SEVERITY CLASSIFICATION & RESOLUTION SLA (UAT & WARRANTY)").font = font_section
    ws_acc.cell(row=r, column=1).fill = fill_section
    ws_acc.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_acc.row_dimensions[r].height = 24
    r += 1
    headers_sla = ["Severity Level", "Definition", "Example Scenario", "Initial Response SLA", "Target Resolution SLA", "Exit Gate Impact"]
    ws_acc.row_dimensions[r].height = 24
    for i, h in enumerate(headers_sla, 1):
        cell = ws_acc.cell(row=r, column=i, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    r += 1
    for s_row in p_data["sla_data"]:
        ws_acc.row_dimensions[r].height = 32
        for ci, val in enumerate(s_row, 1):
            cell = ws_acc.cell(row=r, column=ci, value=val)
            cell.font = font_cell
            cell.fill = fill_zebra if r % 2 == 0 else fill_white
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if ci in [1, 4, 5] else "left", vertical="center", wrap_text=True)
        r += 1
    auto_fit(ws_acc, {"A": 16, "B": 28, "C": 35, "D": 18, "E": 18, "F": 30})

    # 6. FR
    ws_fr = wb.create_sheet(title="FR")
    ws_fr.merge_cells("A1:G1")
    ws_fr["A1"] = "FUNCTIONAL REQUIREMENTS MATRIX (FR)"
    ws_fr["A1"].font = font_title
    ws_fr.row_dimensions[1].height = 28
    ws_fr.merge_cells("A2:G2")
    ws_fr["A2"] = "Detailed functional breakdown across platform modules."
    ws_fr["A2"].font = font_subtitle
    ws_fr.row_dimensions[2].height = 18
    headers_fr = ["Requirement ID", "Functional Module", "Feature Name", "Detailed User Story & Specification", "Priority (MoSCoW)", "Complexity", "Sprint Allocation"]
    ws_fr.row_dimensions[4].height = 24
    for i, h in enumerate(headers_fr, 1):
        cell = ws_fr.cell(row=4, column=i, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_idx, r_data in enumerate(p_data["fr_data"], 5):
        ws_fr.row_dimensions[r_idx].height = 26
        fill_curr = fill_zebra if r_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(r_data, 1):
            cell = ws_fr.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.fill = fill_curr
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [1, 2, 5, 6, 7] else "left", vertical="center", wrap_text=True)
    auto_fit(ws_fr, {"A": 16, "B": 22, "C": 30, "D": 55, "E": 18, "F": 14, "G": 16})

    # 7. Tech-stack
    ws_tech = wb.create_sheet(title="Tech-stack")
    ws_tech.merge_cells("A1:F1")
    ws_tech["A1"] = "TECHNICAL ARCHITECTURE & TECHNOLOGY STACK"
    ws_tech["A1"].font = font_title
    ws_tech.row_dimensions[1].height = 28
    ws_tech.merge_cells("A2:F2")
    ws_tech["A2"] = "Tiered architecture, technology choices, and architectural rationale."
    ws_tech["A2"].font = font_subtitle
    ws_tech.row_dimensions[2].height = 18
    headers_tech = ["Architecture Tier", "Component / Framework", "Version", "Selection Rationale & Technical Justification", "Licensing & Hosting", "High Availability"]
    ws_tech.row_dimensions[4].height = 24
    for i, h in enumerate(headers_tech, 1):
        cell = ws_tech.cell(row=4, column=i, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_idx, r_data in enumerate(p_data["tech_data"], 5):
        ws_tech.row_dimensions[r_idx].height = 28
        fill_curr = fill_zebra if r_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(r_data, 1):
            cell = ws_tech.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.fill = fill_curr
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [1, 3, 5, 6] else "left", vertical="center", wrap_text=True)
    auto_fit(ws_tech, {"A": 22, "B": 28, "C": 12, "D": 45, "E": 24, "F": 22})

    # 8. Non-FR
    ws_nfr = wb.create_sheet(title="Non-FR")
    ws_nfr.merge_cells("A1:F1")
    ws_nfr["A1"] = "NON-FUNCTIONAL REQUIREMENTS (NFR) & ARCHITECTURAL BASELINE"
    ws_nfr["A1"].font = font_title
    ws_nfr.row_dimensions[1].height = 28
    ws_nfr.merge_cells("A2:F2")
    ws_nfr["A2"] = "Performance SLAs, system availability, data security, and disaster recovery."
    ws_nfr["A2"].font = font_subtitle
    ws_nfr.row_dimensions[2].height = 18
    headers_nfr = ["NFR ID", "Quality Attribute", "Target Metric / Specification", "Measurement Methodology", "Architecture Enforcement Mechanism", "Verification Phase"]
    ws_nfr.row_dimensions[4].height = 24
    for i, h in enumerate(headers_nfr, 1):
        cell = ws_nfr.cell(row=4, column=i, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_idx, r_data in enumerate(p_data["nfr_data"], 5):
        ws_nfr.row_dimensions[r_idx].height = 28
        fill_curr = fill_zebra if r_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(r_data, 1):
            cell = ws_nfr.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.fill = fill_curr
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [1, 2, 6] else "left", vertical="center", wrap_text=True)
    auto_fit(ws_nfr, {"A": 16, "B": 22, "C": 35, "D": 30, "E": 40, "F": 20})

    # 9. RateCard
    ws_rate = wb.create_sheet(title="RateCard")
    ws_rate.merge_cells("A1:F1")
    ws_rate["A1"] = "STANDARD PROFESSIONAL SERVICES RATE CARD"
    ws_rate["A1"].font = font_title
    ws_rate.row_dimensions[1].height = 28
    ws_rate.merge_cells("A2:F2")
    ws_rate["A2"] = "Standard hourly and daily rates by seniority and role (USD ex Tax)."
    ws_rate["A2"].font = font_subtitle
    ws_rate.row_dimensions[2].height = 18
    headers_rate = ["Role Title", "Seniority Level", "Hourly Rate (USD)", "Daily Rate (USD)", "Weekly Rate (USD)", "Typical Experience Profile"]
    ws_rate.row_dimensions[4].height = 24
    for i, h in enumerate(headers_rate, 1):
        cell = ws_rate.cell(row=4, column=i, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_idx, r_data in enumerate(p_data["rate_data"], 5):
        ws_rate.row_dimensions[r_idx].height = 24
        fill_curr = fill_zebra if r_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(r_data, 1):
            cell = ws_rate.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.fill = fill_curr
            cell.border = thin_border
            if c_idx in [3, 4, 5]:
                cell.number_format = '$#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx == 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    auto_fit(ws_rate, {"A": 26, "B": 16, "C": 18, "D": 18, "E": 18, "F": 45})

    # 10. Macro Plan & Quotation
    ws_quote = wb.create_sheet(title="Macro Plan & Quotation")
    ws_quote.merge_cells("A1:G1")
    ws_quote["A1"] = "MACRO IMPLEMENTATION PLAN & DETAILED COMMERCIAL QUOTATION"
    ws_quote["A1"].font = font_title
    ws_quote.row_dimensions[1].height = 28
    ws_quote.merge_cells("A2:G2")
    ws_quote["A2"] = "Phased timeline, Work Breakdown Structure (WBS), Quotation, and Milestone Gates."
    ws_quote["A2"].font = font_subtitle
    ws_quote.row_dimensions[2].height = 18

    ws_quote.merge_cells("A4:G4")
    ws_quote["A4"] = f"1. PHASED MACRO ROADMAP & DELIVERABLE MILESTONES ({p_data['duration_weeks']} WEEKS)"
    ws_quote["A4"].font = font_section
    ws_quote["A4"].fill = fill_section
    ws_quote["A4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_quote.row_dimensions[4].height = 24

    headers_phases = ["Phase #", "Phase Name & Objective", "Start (Wk)", "Duration", "Key Activities & Milestones", "Primary Deliverables", "Gate Sign-off"]
    ws_quote.row_dimensions[5].height = 24
    for i, h in enumerate(headers_phases, 1):
        cell = ws_quote.cell(row=5, column=i, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    r = 6
    for p_row in p_data["phases_data"]:
        ws_quote.row_dimensions[r].height = 30
        for ci, val in enumerate(p_row, 1):
            cell = ws_quote.cell(row=r, column=ci, value=val)
            cell.font = font_cell
            cell.fill = fill_zebra if r % 2 == 0 else fill_white
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if ci in [1, 3, 4, 7] else "left", vertical="center", wrap_text=True)
        r += 1

    r += 1
    ws_quote.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws_quote.cell(row=r, column=1, value="2. WORK BREAKDOWN STRUCTURE (WBS) & EFFORT ESTIMATION").font = font_section
    ws_quote.cell(row=r, column=1).fill = fill_section
    ws_quote.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_quote.row_dimensions[r].height = 24
    r += 1

    headers_wbs = ["WBS ID", "Deliverable Component", "Est. Man-Days", "Blended Day Rate", "Subtotal (USD)", "Resource Composition", "Billing Type"]
    ws_quote.row_dimensions[r].height = 24
    for i, h in enumerate(headers_wbs, 1):
        cell = ws_quote.cell(row=r, column=i, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    r += 1

    subtotal_start_row = r
    for w_row in p_data["wbs_data"]:
        ws_quote.row_dimensions[r].height = 24
        for ci, val in enumerate(w_row, 1):
            cell = ws_quote.cell(row=r, column=ci, value=val)
            cell.font = font_cell
            cell.fill = fill_zebra if r % 2 == 0 else fill_white
            cell.border = thin_border
            if ci in [4, 5]:
                cell.number_format = '$#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif ci in [1, 3, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        r += 1

    totals = [
        (f"Subtotal Professional Services ({p_data['total_mandays']} Man-Days)", f"=SUM(E{subtotal_start_row}:E{r-1})", font_cell_bold),
        ("Project Delivery Contingency Buffer (10%)", f"=E{r}*0.1", font_cell),
        ("TOTAL DELIVERY QUOTATION (USD ex GST/Tax)", f"=E{r}+E{r+1}", font_title)
    ]
    for label, formula, f_style in totals:
        ws_quote.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        lbl_c = ws_quote.cell(row=r, column=1, value=label)
        lbl_c.font = f_style
        lbl_c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        for ci in range(1, 5):
            ws_quote.cell(row=r, column=ci).border = thin_border
            ws_quote.cell(row=r, column=ci).fill = fill_sub if "TOTAL" in label else fill_white

        val_c = ws_quote.cell(row=r, column=5, value=formula)
        val_c.font = f_style
        val_c.fill = fill_sub if "TOTAL" in label else fill_white
        val_c.border = thin_border
        val_c.number_format = '$#,##0.00'
        val_c.alignment = Alignment(horizontal="right", vertical="center")

        for ci in range(6, 8):
            ws_quote.cell(row=r, column=ci).border = thin_border
            ws_quote.cell(row=r, column=ci).fill = fill_sub if "TOTAL" in label else fill_white
        ws_quote.row_dimensions[r].height = 26
        r += 1

    r += 1
    ws_quote.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws_quote.cell(row=r, column=1, value="3. DELIVERABLE-BASED MILESTONE PAYMENT SCHEDULE").font = font_section
    ws_quote.cell(row=r, column=1).fill = fill_section
    ws_quote.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_quote.row_dimensions[r].height = 24
    r += 1

    headers_pay = ["Milestone #", "Payment Trigger & Deliverable Gate", "% Allocation", "Amount (USD)", "Project Target Date", "Invoice Terms"]
    ws_quote.row_dimensions[r].height = 24
    for i, h in enumerate(headers_pay, 1):
        cell = ws_quote.cell(row=r, column=i, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    r += 1

    for pay_row in p_data["payments_data"]:
        ws_quote.row_dimensions[r].height = 24
        for ci, val in enumerate(pay_row, 1):
            cell = ws_quote.cell(row=r, column=ci, value=val)
            cell.font = font_cell
            cell.fill = fill_zebra if r % 2 == 0 else fill_white
            cell.border = thin_border
            if ci == 3:
                cell.number_format = '0.0%'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif ci == 4:
                cell.number_format = '$#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif ci in [1, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        r += 1

    auto_fit(ws_quote, {"A": 16, "B": 35, "C": 15, "D": 20, "E": 22, "F": 28, "G": 20})
    wb.save(file_path)
    print(f"✅ Generated Proposal: {file_path}")

